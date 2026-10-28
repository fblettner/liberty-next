"""``/api/explorer`` + the admin round trip — a SQLite model with a fact and a dimension."""

from __future__ import annotations

import asyncio
import textwrap
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine

from liberty.auth.db import AuthDatabase
from liberty.auth.service import AuthService
from liberty.config import AISettings, AppSettings, AuthSettings, ConnectorSettings, ExplorerSettings, ScreenSettings, Settings
from liberty.connectors.config import PoolConfig
from liberty.connectors.db import PoolRegistry
from liberty.main import create_app

JWT_SECRET = "web-explorer-test-secret"


def _connectors_toml(db_url: str) -> str:
    return textwrap.dedent(f"""
        [pools.default]
        url = "{db_url}"

        [connectors.shop]
        type = "sql"
        pool = "default"
        [[connectors.shop.queries]]
        name = "fact_sales_get"
        sql = "SELECT id AS ID, company AS COMPANY, sale_date AS SALE_DATE, qty AS QTY, revenue AS REVENUE FROM fact_sales"
        [[connectors.shop.queries]]
        name = "dim_company_get"
        sql = "SELECT company AS COMPANY, name AS NAME, region AS REGION FROM dim_company"
        [[connectors.shop.queries]]
        name = "broken_get"
        sql = "SELECT nope FROM missing_table"
        """)


EXPLORER_TOML = textwrap.dedent("""
    [explorer.shop.sales]
    label = "Shop"
    hub = "sales"

    [explorer.shop.sales.entities.sales]
    query = "fact_sales_get"
    screen = "sales"
    key = ["ID"]
    default_date = "SALE_DATE"
    measures = { QTY = "sum", REVENUE = "sum", COMPANY = "sum" }
    hidden = ["ID"]

    [explorer.shop.sales.entities.company]
    label = "Company"
    query = "dim_company_get"
    key = ["COMPANY"]

    [[explorer.shop.sales.relations]]
    from = "sales"
    to = "company"
    on = [["COMPANY", "COMPANY"]]

    [explorer.other.x]
    hub = "x"
    [explorer.other.x.entities.x]
    connector = "shop"
    query = "broken_get"

    [explorer.shop.frompreset]           # a preset named, not applied yet (the builder fills it)
    label = "From the preset"
    preset = "shop_sales"
    """)

PRESETS_TOML = textwrap.dedent("""
    [[presets]]
    id = "shop_sales"
    label = "Sales + company"
    connector = "shop"
    [[presets.tables]]
    name = "fact_sales"
    key = ["id"]
    [[presets.tables]]
    name = "dim_company"
    join = "LEFT"
    on = [{ left = "fact_sales.COMPANY", right = "dim_company.COMPANY" }]
    key = ["COMPANY", "NOPE"]

    [[presets]]                          # a preset that ships its model (explorer.shop.shop_shipped)
    id = "shop_shipped"
    label = "Sales (shipped)"
    description = "The model the preset ships."
    connector = "shop"
    app = "shop"
    explorer = true
    [[presets.tables]]
    name = "fact_sales"
    label = "Sales lines"
    key = ["id"]
    [[presets.tables]]
    name = "dim_company"
    label = "Company (dim)"
    cardinality = "one"
    join = "LEFT"
    on = [{ left = "fact_sales.COMPANY", right = "dim_company.COMPANY" }]
    key = ["COMPANY"]
    """)

SCREENS_TOML = textwrap.dedent("""
    [screens.shop.sales]
    label = "Sales"
    read_query = "fact_sales_get"
    [[screens.shop.sales.columns]]
    name = "QTY"
    label = "Quantity"
    [[screens.shop.sales.columns]]
    name = "SALE_DATE"
    format = "date"
    """)


def _seed(db_url: str) -> None:
    async def go() -> None:
        engine = create_async_engine(db_url)
        async with engine.begin() as c:
            await c.execute(text("CREATE TABLE fact_sales (id INTEGER PRIMARY KEY, company TEXT, sale_date DATE, qty INTEGER, revenue NUMERIC)"))
            await c.execute(text("CREATE TABLE dim_company (company TEXT PRIMARY KEY, name TEXT, region TEXT)"))
            await c.execute(text("INSERT INTO fact_sales VALUES (1, 'C1', '2026-01-05', 3, 30.5)"))
            await c.execute(text("INSERT INTO dim_company VALUES ('C1', 'Corp one', 'EU')"))
        await engine.dispose()
        pools = PoolRegistry({"default": PoolConfig(url=db_url)})
        db = AuthDatabase(pools, "default")
        await db.create_schema()
        async with db.session() as s:
            svc = AuthService(s)
            await svc.get_or_create_role("admin", permissions=["*"])
            await svc.get_or_create_role("analyst", permissions=["explorer:shop:*", "sql:shop:*"])
            await svc.create_user("admin", password="adminpw", is_superuser=True, roles=["admin"])
            await svc.create_user("analyst", password="analystpw", roles=["analyst"])
            await svc.create_user("nobody", password="nobodypw")
        await pools.dispose()

    asyncio.run(go())


@pytest.fixture
def env(tmp_path):
    db_url = f"sqlite+aiosqlite:///{tmp_path / 'app.db'}"
    (tmp_path / "connectors.toml").write_text(_connectors_toml(db_url))
    (tmp_path / "explorer.toml").write_text(EXPLORER_TOML)
    (tmp_path / "screens.toml").write_text(SCREENS_TOML)
    (tmp_path / "presets").mkdir()
    (tmp_path / "presets" / "shop.toml").write_text(PRESETS_TOML)
    _seed(db_url)
    settings = Settings(
        app=AppSettings(static_dir=""),
        connectors=ConnectorSettings(config_path=Path(tmp_path / "connectors.toml")),
        screens=ScreenSettings(config_path=tmp_path / "screens.toml"),
        explorer=ExplorerSettings(config_path=tmp_path / "explorer.toml"),
        auth=AuthSettings(backend="db", jwt_secret=JWT_SECRET, pool="default"),
        ai=AISettings(enabled=False),
    )
    return create_app(settings), tmp_path


def _h(client: TestClient, username: str) -> dict[str, str]:
    tok = client.post("/auth/login", json={"username": username, "password": f"{username}pw"}).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def test_list_models_by_permission(env) -> None:
    app, _ = env
    with TestClient(app) as client:
        assert client.get("/api/explorer").status_code == 401
        admin = client.get("/api/explorer", headers=_h(client, "admin")).json()["models"]
        assert [(m["app"], m["model"]) for m in admin] == [("shop", "sales"), ("shop", "frompreset"), ("shop", "shop_shipped"), ("other", "x")] and admin[0]["label"] == "Shop" and admin[0]["entities"] == 2
        assert admin[2]["derived"] is True and admin[2]["preset"] == "shop_shipped" and admin[0]["derived"] is False
        assert [m["model"] for m in client.get("/api/explorer", headers=_h(client, "analyst")).json()["models"]] == ["sales", "frompreset", "shop_shipped"]
        assert client.get("/api/explorer", headers=_h(client, "nobody")).json()["models"] == []


def test_model_fields_discovered_with_hints(env) -> None:
    app, _ = env
    with TestClient(app) as client:
        r = client.get("/api/explorer/shop/sales/model", headers=_h(client, "analyst"))
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["hub"] == "sales" and body["relations"] == [{"from": "sales", "to": "company", "on": [["COMPANY", "COMPANY"]]}]
        sales, company = body["entities"]
        assert sales["id"] == "sales" and sales["cardinality"] == "one" and sales["default_date"] == "SALE_DATE"
        by = {f["name"]: f for f in sales["fields"]}
        assert "ID" not in by                                     # hidden
        assert by["QTY"] == {"key": "sales.QTY", "entity": "sales", "name": "QTY", "label": "Quantity", "type": "number", "kind": "measure", "aggregate": "sum"}
        assert by["SALE_DATE"]["type"] == "date" and by["SALE_DATE"]["format"] == "date"
        assert by["COMPANY"]["kind"] == "dimension"               # declared a measure, but text
        assert by["REVENUE"]["kind"] == "measure" and by["REVENUE"]["type"] == "number"
        assert company["label"] == "Company" and [f["key"] for f in company["fields"]] == ["company.COMPANY", "company.NAME", "company.REGION"]
        assert any("COMPANY" in w and "typed string" in w for w in body["warnings"])
        # permission: nobody → 403; unknown app (with permission) → 404
        assert client.get("/api/explorer/shop/sales/model", headers=_h(client, "nobody")).status_code == 403
        assert client.get("/api/explorer/ghost/x/model", headers=_h(client, "admin")).status_code == 404


def test_model_with_a_broken_entity_reports_it(env) -> None:
    app, _ = env
    with TestClient(app) as client:
        body = client.get("/api/explorer/other/x/model", headers=_h(client, "admin")).json()
        assert body["entities"][0]["fields"] == [] and "broken_get" in body["entities"][0]["error"]
        assert body["warnings"] and "broken_get" in body["warnings"][0]


def test_model_cache_and_refresh(env) -> None:
    app, tmp_path = env
    with TestClient(app) as client:
        h = _h(client, "admin")
        first = client.get("/api/explorer/shop/sales/model", headers=h).json()
        assert app.state.explorer_cache._views  # cached
        # editing the file + reload empties the cache and the new label is seen
        (tmp_path / "explorer.toml").write_text(EXPLORER_TOML.replace('label = "Shop"', 'label = "Shop v2"'))
        client.post("/admin/reload", headers=h)
        second = client.get("/api/explorer/shop/sales/model", headers=h).json()
        assert first["label"] == "Shop" and second["label"] == "Shop v2"
        assert client.get("/api/explorer/shop/sales/model?refresh=1", headers=h).json()["label"] == "Shop v2" and first["model"] == "sales"


def test_admin_parsed_round_trip(env) -> None:
    app, tmp_path = env
    with TestClient(app) as client:
        h = _h(client, "admin")
        body = client.get("/admin/config/explorer/parsed", headers=h).json()
        shop = body["explorer"]["shop"]["sales"]
        assert shop["hub"] == "sales" and shop["entities"]["sales"]["measures"]["QTY"] == "sum"
        assert shop["relations"][0]["from"] == "sales" and "id" not in shop["entities"]["sales"]
        shop["entities"]["company"]["label"] = "Companies"
        shop["entities"]["region"] = {"query": "dim_company_get", "key": ["REGION"]}
        shop["relations"].append({"from": "company", "to": "region", "on": [["REGION", "REGION"]]})
        r = client.put("/admin/config/explorer/parsed", json={"explorer": body["explorer"]}, headers=h)
        assert r.status_code == 200, r.text
        text_ = (tmp_path / "explorer.toml").read_text()
        assert 'label = "Companies"' in text_ and 'from = "company"' in text_
        client.post("/admin/reload", headers=h)
        back = client.get("/admin/config/explorer/parsed", headers=h).json()["explorer"]["shop"]["sales"]
        assert back["entities"]["company"]["label"] == "Companies" and len(back["relations"]) == 2
        # an invalid model is refused and the file untouched
        bad = {"explorer": {"shop": {"sales": {**shop, "hub": "ghost"}}}}
        assert client.put("/admin/config/explorer/parsed", json=bad, headers=h).status_code == 422
        assert client.get("/admin/config/explorer/parsed", headers=_h(client, "analyst")).status_code == 403
        assert "explorer" in client.get("/admin/config/schema", headers=h).json()
        assert client.get("/info", headers=h).json()["explorer"] == {"apps": ["shop", "other"], "models": 4}   # + the preset-shipped one


def test_query_route(env) -> None:
    app, _ = env
    with TestClient(app) as client:
        h = _h(client, "analyst")
        body = {"rows": [{"field": "company.NAME"}], "measures": [{"field": "sales.REVENUE"}], "debug": True}
        r = client.post("/api/explorer/shop/sales/query", json=body, headers=h)
        assert r.status_code == 200, r.text
        out = r.json()
        assert out["rows"] == [{"company.NAME": "Corp one", "sum(sales.REVENUE)": 30.5}] and out["groups"] == 1
        assert "sql" not in out                                  # debug SQL is for a superuser
        assert "sql" in client.post("/api/explorer/shop/sales/query", json=body, headers=_h(client, "admin")).json()
        assert client.post("/api/explorer/shop/sales/query", json={"rows": [{"field": "sales.NOPE"}]}, headers=h).status_code == 422
        assert client.post("/api/explorer/shop/sales/query", json={"rows": [{"field": "sales.COMPANY", "bucket": "month"}]}, headers=h).status_code == 422
        assert client.post("/api/explorer/shop/sales/query", json=body, headers=_h(client, "nobody")).status_code == 403
        assert client.post("/api/explorer/ghost/x/query", json=body, headers=_h(client, "admin")).status_code == 404
        # a broken entity: a database error, reported with its text to a superuser
        r = client.post("/api/explorer/other/x/query", json={}, headers=_h(client, "admin"))
        assert r.status_code == 502 and "missing_table" in r.text


def test_detail_and_calc_routes(env) -> None:
    app, _ = env
    with TestClient(app) as client:
        h = _h(client, "analyst")
        r = client.post("/api/explorer/shop/sales/detail", json={"columns": ["sales.ID", "company.NAME"], "drill": [{"field": "company.NAME", "value": "Corp one"}]}, headers=h)
        assert r.status_code == 200, r.text
        assert r.json()["total"] == 1 and r.json()["rows"] == [{"sales.ID": 1, "company.NAME": "Corp one"}]
        r = client.post("/api/explorer/shop/sales/calc/test", json={"calculated": [{"name": "up", "expr": "REVENUE / QTY"}], "name": "up"}, headers=h)
        assert r.status_code == 200, r.text
        assert r.json() == {"name": "up", "type": "number", "kind": "measure", "values": [10.166666666666666]}
        r = client.post("/api/explorer/shop/sales/calc/test", json={"calculated": [{"name": "up", "expr": "REVENUE +"}], "name": "up"}, headers=h)
        assert r.status_code == 422 and "unexpected" in r.json()["detail"]
        assert client.post("/api/explorer/shop/sales/detail", json={"columns": ["sales.NOPE"]}, headers=h).status_code == 422
        assert client.post("/api/explorer/shop/sales/detail", json={}, headers=_h(client, "nobody")).status_code == 403


def test_saved_analyses_routes(env) -> None:
    app, _ = env
    with TestClient(app) as client:
        ha, hb = _h(client, "analyst"), _h(client, "admin")
        r = client.put("/api/explorer/shop/sales/saved/analyses", json={"name": "Revenue by company", "definition": {"rows": [{"field": "company.NAME"}]}, "audience": "sales"}, headers=ha)
        assert r.status_code == 200, r.text
        item = r.json()["item"]
        assert item["id"] == "revenue-by-company" and item["owner"] == "analyst" and not item["shared"]
        assert client.get("/api/explorer/shop/sales/saved/analyses", headers=hb).json()["items"] == []     # private to the analyst
        r = client.put("/api/explorer/shop/sales/saved/analyses", json={"id": item["id"], "name": "Revenue by company", "definition": {}, "shared": True}, headers=ha)
        assert r.json()["item"]["shared"] and r.json()["item"]["version"] == 2
        assert [x["id"] for x in client.get("/api/explorer/shop/sales/saved/analyses", headers=hb).json()["items"]] == ["revenue-by-company"]
        # the admin cannot be refused; nobody has no permission; a bad kind is a 404
        assert client.get("/api/explorer/shop/sales/saved/analyses", headers=_h(client, "nobody")).status_code == 403
        assert client.get("/api/explorer/shop/sales/saved/things", headers=ha).status_code == 404
        r = client.put("/api/explorer/shop/sales/saved/dashboards", json={"name": "Home", "definition": {"tiles": [{"analysis": "revenue-by-company", "w": 2, "h": "m"}]}, "shared": True}, headers=hb)
        assert r.status_code == 200 and r.json()["item"]["id"] == "home"
        assert client.delete("/api/explorer/shop/sales/saved/dashboards/home", headers=ha).status_code == 403   # shared by the admin
        assert client.delete("/api/explorer/shop/sales/saved/dashboards/home", headers=hb).json()["deleted"] is True
        assert client.delete("/api/explorer/shop/sales/saved/analyses/revenue-by-company", headers=ha).json()["deleted"] is True


def test_export_routes(env) -> None:
    import io
    from openpyxl import load_workbook
    app, _ = env
    with TestClient(app) as client:
        h = _h(client, "analyst")
        definition = {"entity": "sales", "view": "table", "rows": [{"field": "company.NAME"}], "measures": [{"field": "sales.REVENUE", "agg": "sum"}], "filterRows": [], "sort": [], "limit": 100, "calculated": []}
        r = client.post("/api/explorer/shop/sales/export", json={"format": "xlsx", "definition": definition, "title": "Revenue by company"}, headers=h)
        assert r.status_code == 200, r.text
        assert r.headers["content-disposition"] == 'attachment; filename="Revenue_by_company.xlsx"'
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Revenue by company"] and wb.active["A3"].value == "Corp one" and wb.active["B3"].value == 30.5   # row 1 = the note
        # a saved analysis + a dashboard of it, by id
        client.put("/api/explorer/shop/sales/saved/analyses", json={"name": "Rev", "definition": definition, "shared": True}, headers=h)
        client.put("/api/explorer/shop/sales/saved/dashboards", json={"name": "Board", "definition": {"tiles": [{"analysis": "rev", "w": 2, "h": "m"}, {"analysis": "ghost", "w": 1, "h": "s"}]}, "shared": True}, headers=h)
        r = client.post("/api/explorer/shop/sales/export", json={"format": "xlsx", "dashboard": "board"}, headers=h)
        assert r.status_code == 200 and load_workbook(io.BytesIO(r.content)).sheetnames == ["Rev"]
        assert client.post("/api/explorer/shop/sales/export", json={"format": "xlsx", "analysis": "nope"}, headers=h).status_code == 404
        assert client.post("/api/explorer/shop/sales/export", json={"format": "xlsx"}, headers=h).status_code == 422
        assert client.post("/api/explorer/shop/sales/export", json={"format": "xlsx", "analysis": "rev"}, headers=_h(client, "nobody")).status_code == 403
        # pdf: the reports extra may be missing on this machine
        r = client.post("/api/explorer/shop/sales/export", json={"format": "pdf", "analysis": "rev", "from": "2026-01-01", "to": "2026-12-31"}, headers=h)
        assert r.status_code in (200, 503)
        if r.status_code == 200:
            assert r.content[:4] == b"%PDF"


def test_model_from_a_preset(env) -> None:
    """Settings → Explorer applies a preset: the endpoint returns the model filled — the preset's
    tables as entities (lower-cased ids, SELECT * on each), the base table as hub, the join as a
    relation resolved on the discovered columns, the declared entity completed — and the model
    saved that way runs. A model that names a preset but has no entity is built from the preset
    at first use (the preset's warnings on the view). Applying the preset also creates the
    dictionary entries the tables lack (a non-JDE source: the inferred format, no label)."""
    app, tmp_path = env
    with TestClient(app) as client:
        h = _h(client, "admin")
        body = client.get("/api/explorer/shop/frompreset/model", headers=h).json()
        assert body["hub"] == "fact_sales" and [e["id"] for e in body["entities"]] == ["fact_sales", "dim_company"] and "NOPE" in body["warnings"][0]
        assert body["derived"] is False and body["undocumented"] == 8
        r = client.post("/admin/explorer/from-preset", json={"app": "shop", "preset": "shop_sales", "model": {"entities": {"fact_sales": {"table": "fact_sales", "measures": {"QTY": "sum", "REVENUE": "sum"}}}}}, headers=h)
        assert r.status_code == 200, r.text
        filled = r.json()
        assert filled["warnings"] == ["preset key of dim_company: column not found for NOPE"] and filled["model"]["hub"] == "fact_sales" and filled["model"]["preset"] == "shop_sales"
        assert filled["dictionary"] == {"created": ["COMPANY", "ID", "NAME", "QTY", "REGION", "REVENUE", "SALE_DATE"], "scanned": 2, "skipped": []}
        assert "[connectors.shop.entries.SALE_DATE]" in (tmp_path / "dictionary.toml").read_text()
        # a second apply creates nothing (every entry exists now); the fields are documented
        again = client.post("/admin/explorer/from-preset", json={"app": "shop", "preset": "shop_sales"}, headers=h).json()
        assert again["dictionary"]["created"] == []
        body = client.get("/api/explorer/shop/frompreset/model", headers=h).json()
        assert body["undocumented"] == 0 and {f["name"]: f.get("format") for f in body["entities"][0]["fields"]}["SALE_DATE"] == "date"
        ents = filled["model"]["entities"]
        # the preset's keys resolved to the real columns (`id` → ID); a missing item is only a warning
        assert set(ents) == {"fact_sales", "dim_company"} and ents["dim_company"] == {"label": "dim_company", "connector": "shop", "table": "dim_company", "key": ["COMPANY"], "cardinality": "many"}
        assert ents["fact_sales"]["key"] == ["ID"] and ents["fact_sales"]["connector"] == "shop" and ents["fact_sales"]["table"] == "fact_sales"
        assert filled["model"]["relations"] == [{"from": "dim_company", "to": "fact_sales", "on": [["COMPANY", "COMPANY"]]}]
        assert client.post("/admin/explorer/from-preset", json={"app": "shop", "preset": "ghost"}, headers=h).status_code == 404
        assert client.post("/admin/explorer/from-preset", json={"app": "shop", "preset": "shop_sales"}, headers=_h(client, "analyst")).status_code == 403
        # a table entity's columns for the editor
        d = client.post("/admin/explorer/describe-entity", json={"app": "shop", "entity": {"table": "dim_company", "connector": "shop"}}, headers=h).json()
        assert [c["name"] for c in d["columns"]] == ["COMPANY", "NAME", "REGION"] and d["error"] is None
        # saved as returned, the model runs
        doc = client.get("/admin/config/explorer/parsed", headers=h).json()["explorer"]
        doc["shop"]["frompreset"] = filled["model"]
        assert client.put("/admin/config/explorer/parsed", json={"explorer": doc}, headers=h).status_code == 200
        client.post("/admin/reload", headers=h)
        body = client.get("/api/explorer/shop/frompreset/model", headers=h).json()
        assert body["hub"] == "fact_sales" and body["warnings"] == [] and {f["name"]: f["kind"] for f in body["entities"][0]["fields"]}["QTY"] == "measure"
        r = client.post("/api/explorer/shop/frompreset/query", json={"rows": [{"field": "dim_company.NAME"}], "measures": [{"field": "fact_sales.REVENUE"}]}, headers=h)
        assert r.status_code == 200, r.text
        assert r.json()["rows"] == [{"dim_company.NAME": "Corp one", "sum(fact_sales.REVENUE)": 30.5}]


def test_preset_shipped_model_and_its_dictionary(env) -> None:
    """A preset flagged ``explorer = true`` with an ``app`` is a model by default — listed, built
    from the preset at first use (the table labels and cardinalities of the preset), queryable —
    and stays out of explorer.toml. The page's "create the dictionary entries" scans its tables."""
    app, _ = env
    with TestClient(app) as client:
        h = _h(client, "admin")
        body = client.get("/api/explorer/shop/shop_shipped/model", headers=h)
        assert body.status_code == 200, body.text
        body = body.json()
        assert body["derived"] is True and body["preset"] == "shop_shipped" and body["label"] == "Sales (shipped)" and body["hub"] == "fact_sales"
        sales, company = body["entities"]
        assert sales["label"] == "Sales lines" and sales["key"] == ["ID"] and company["label"] == "Company (dim)" and company["cardinality"] == "one"
        assert body["relations"] == [{"from": "dim_company", "to": "fact_sales", "on": [["COMPANY", "COMPANY"]]}] and body["warnings"] == []
        assert body["undocumented"] == 8
        r = client.post("/api/explorer/shop/shop_shipped/query", json={"rows": [{"field": "dim_company.NAME"}], "measures": [{"field": "fact_sales.REVENUE"}]}, headers=h)
        assert r.status_code == 200, r.text
        assert r.json()["rows"] == [{"dim_company.NAME": "Corp one", "sum(fact_sales.REVENUE)": 30.5}]
        # not in the file (the builder lists what explorer.toml declares)
        assert "shop_shipped" not in client.get("/admin/config/explorer/parsed", headers=h).json()["explorer"]["shop"]
        # the dictionary entries — superuser only, unknown model 404
        assert client.post("/admin/explorer/shop/shop_shipped/dictionary", headers=_h(client, "analyst")).status_code == 403
        assert client.post("/admin/explorer/shop/ghost/dictionary", headers=h).status_code == 404
        r = client.post("/admin/explorer/shop/shop_shipped/dictionary", headers=h)
        assert r.status_code == 200, r.text
        assert r.json() == {"created": ["COMPANY", "ID", "NAME", "QTY", "REGION", "REVENUE", "SALE_DATE"], "scanned": 2, "skipped": []}
        body = client.get("/api/explorer/shop/shop_shipped/model", headers=h).json()
        assert body["undocumented"] == 0 and body["derived"] is True
