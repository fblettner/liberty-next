"""``file_export`` — a named read query to an xlsx / csv file in the run's work area, on a real
SQLite registry: binds from the run parameters and templated ``params``, labels as headers,
templated file names, the copy to ``export_dir``, the refusals."""

from __future__ import annotations

import csv
from types import SimpleNamespace

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import text

from liberty.connectors.config import ConnectorsFile, PoolConfig, QueryDef, SqlConnectorConfig
from liberty.connectors.registry import ConnectorRegistry
from liberty.jobs import FileExportExecutor, ManualTrigger, RunContext, Step, StepFailed


@pytest_asyncio.fixture
async def registry(tmp_path):
    cfg = ConnectorsFile(
        pools={"db": PoolConfig(url=f"sqlite+aiosqlite:///{tmp_path / 'db.sqlite'}")},
        connectors={"sales": SqlConnectorConfig(type="sql", pool="db", queries=[
            QueryDef(name="sales_get", sql="SELECT company, qty, revenue FROM sales WHERE period = :PERIOD ORDER BY company"),
            QueryDef(name="sales_put", sql="UPDATE sales SET qty = 0", writable=True),
            QueryDef(name="broken_get", sql="SELECT * FROM no_such_table"),
        ])},
    )
    reg = ConnectorRegistry(cfg)
    async with reg.pools.engine("db").begin() as conn:
        await conn.execute(text("CREATE TABLE sales (period TEXT, company TEXT, qty INTEGER, revenue NUMERIC(14, 2))"))
        await conn.execute(text("INSERT INTO sales VALUES ('2026-03', 'B', 2, 20.5), ('2026-03', 'A', 1, 10), ('2026-02', 'C', 9, 90)"))
    try:
        yield reg
    finally:
        await reg.aclose()


def _settings(tmp_path):
    return SimpleNamespace(jobs=SimpleNamespace(files_dir=str(tmp_path / "files"), config_path=str(tmp_path / "jobs.toml")))


def _ctx() -> RunContext:
    return RunContext(run_id="run-1", job_id="export-job", trigger=ManualTrigger(triggered_by="tester"))


def _step(**kw) -> Step:
    return Step(**{"type": "file_export", "name": "export", "connector": "sales", "query": "sales_get", **kw})


async def test_xlsx_from_the_run_parameters(tmp_path, registry) -> None:
    ex = FileExportExecutor(registry, settings=_settings(tmp_path))
    s = _step(filename="sales-${params.period}", sheet="Sales", export_dir=str(tmp_path / "out" / "${params.period}"))
    s = s.model_copy(update={"op_kwargs": {"period": "2026-03"}})      # the run parameter, bound as :PERIOD
    res = await ex.execute(s, _ctx())
    assert res.rows_affected == 2 and res.extras["file"] == "run-1-sales-2026-03.xlsx" and res.extras["truncated"] is False
    from liberty.jobs.files import job_dir
    out = job_dir(_settings(tmp_path), "export-job") / "run-1-sales-2026-03.xlsx"
    ws = load_workbook(out)["Sales"]
    values = [[c.value for c in r] for r in ws.iter_rows()]
    assert [v.lower() for v in values[0]] == ["company", "qty", "revenue"]
    assert [r[:2] for r in values[1:]] == [["A", 1], ["B", 2]]
    assert (tmp_path / "out" / "2026-03" / "sales-2026-03.xlsx").exists() and res.extras["copied_to"].endswith("sales-2026-03.xlsx")


async def test_csv_with_a_templated_bind(tmp_path, registry) -> None:
    ex = FileExportExecutor(registry, settings=_settings(tmp_path))
    s = _step(format="csv", delimiter=",", params={"PERIOD": "${params.month}"}).model_copy(update={"op_kwargs": {"month": "2026-02"}})
    res = await ex.execute(s, _ctx())
    from liberty.jobs.files import job_dir
    with (job_dir(_settings(tmp_path), "export-job") / "run-1-sales_get.csv").open(encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    assert res.rows_affected == 1 and [r[0].lower() for r in rows[:1]] == ["company"] and rows[1][:2] == ["C", "9"]


async def test_refusals(tmp_path, registry) -> None:
    ex = FileExportExecutor(registry, settings=_settings(tmp_path))
    with pytest.raises(StepFailed, match="writable"):
        await ex.execute(_step(query="sales_put"), _ctx())
    with pytest.raises(StepFailed, match="file_export"):
        await ex.execute(_step(query="nope_get"), _ctx())
    with pytest.raises(StepFailed, match="no_such_table|broken_get"):
        await ex.execute(_step(query="broken_get"), _ctx())
    from liberty.jobs.files import job_dir
    assert not list(job_dir(_settings(tmp_path), "export-job").glob("run-1-broken_get*"))    # no half file left
    with pytest.raises(StepFailed, match="period"):                                          # a template with no value
        await ex.execute(_step(filename="x-${params.period}"), _ctx())


def test_schema_rules() -> None:
    with pytest.raises(ValueError, match="connector, query|missing"):
        Step(type="file_export", name="e")
    with pytest.raises(ValueError, match="xlsx or csv"):
        _step(format="pdf")
    with pytest.raises(ValueError, match="xlsx or pdf"):
        Step(type="report", name="r", report="report:x:y", format="csv")
