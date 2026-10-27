"""The headless side: period presets, the definition → requests, the pivot flattening, the
sections of an analysis and a dashboard, the xlsx / markdown renderers."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from liberty.explorer.report import (
    build_sections, compile_filters, definition_to_aggregate, definition_to_detail, period_range, pivot_rows, sections_to_markdown,
    sections_to_xlsx,
)
from tests.test_explorer_engine import env  # noqa: F401 — the SQLite model fixture

TODAY = date(2026, 10, 27)


def test_period_range_matches_the_frontend_rules() -> None:
    assert period_range("today", TODAY) == (TODAY, TODAY)
    assert period_range("yesterday", TODAY) == (date(2026, 10, 26), TODAY)
    assert period_range("last7", TODAY) == (date(2026, 10, 21), TODAY)
    assert period_range("last30", TODAY) == (date(2026, 9, 28), TODAY)
    assert period_range("thisMonth", TODAY) == (date(2026, 10, 1), TODAY)
    assert period_range("lastMonth", TODAY) == (date(2026, 9, 1), date(2026, 9, 30))
    assert period_range("thisQuarter", TODAY) == (date(2026, 10, 1), TODAY)
    assert period_range("thisYear", TODAY) == (date(2026, 1, 1), TODAY)
    assert period_range("lastYear", TODAY) == (date(2025, 1, 1), date(2025, 12, 31))
    assert period_range("last12", TODAY) == (date(2025, 10, 28), TODAY)
    assert period_range("lastMonth", date(2026, 3, 31)) == (date(2026, 2, 1), date(2026, 2, 28))
    assert period_range("all", TODAY) is None and period_range("nope", TODAY) is None


def test_definition_to_requests() -> None:
    d = {
        "entity": "sales", "view": "table",
        "rows": [{"field": "company.NAME"}, {"field": "sales.SOLD_ON", "bucket": "month"}],
        "measures": [{"field": "sales.QTY", "agg": "sum"}, {"field": None, "agg": "count"}],
        "filterRows": [{"field": "company.NAME", "op": "startsWith", "value": "C"}, {"field": "sales.QTY", "op": "between", "value": "1", "valueTo": "9"},
                       {"field": "sales.ID", "op": "in", "value": "1, 2"}, {"field": "sales.ID", "op": "gt", "value": ""}],
        "period": {"field": "sales.SOLD_ON", "preset": "thisYear", "range": {"dateFrom": "", "dateTo": ""}},
        "sort": [{"key": "sum(sales.QTY)", "dir": "desc"}], "limit": 50, "calculated": [{"name": "up", "expr": "REVENUE / QTY"}],
    }
    req = definition_to_aggregate(d, today=TODAY)
    assert [r.key for r in req.rows] == ["company.NAME", "sales.SOLD_ON@month"] and [m.key for m in req.measures] == ["sum(sales.QTY)", "count(*)"]
    assert req.filters == {"company.NAME": "C", "company.NAME__op": "startsWith", "sales.QTY__from": "1", "sales.QTY__to": "9", "sales.ID": ["1", "2"], "sales.ID__op": "in"}
    assert req.period is not None and req.period.from_ == date(2026, 1, 1) and req.period.to == TODAY
    assert req.limit == 50 and req.calculated[0].name == "up"
    # a forced window (the digest) replaces the preset; no period field → none
    assert definition_to_aggregate(d, window=(date(2026, 9, 1), date(2026, 9, 30))).period.from_ == date(2026, 9, 1)
    assert definition_to_aggregate({**d, "period": None}, window=(date(2026, 9, 1), date(2026, 9, 30))).period is None
    custom = {**d, "period": {"field": "sales.SOLD_ON", "preset": "custom", "range": {"dateFrom": "2026-02-01", "dateTo": ""}}}
    assert definition_to_aggregate(custom).period.from_ == date(2026, 2, 1)
    det = definition_to_detail({**d, "view": "detail"}, today=TODAY)
    assert det.columns == ["company.NAME", "sales.SOLD_ON"] and det.sort == [] and det.page_size == 3000
    assert compile_filters([]) == {}


def test_pivot_rows() -> None:
    rows = [
        {"c": "C1", "m": "2026-01", "v": 3}, {"c": "C1", "m": "2026-02", "v": 2}, {"c": "C2", "m": "2026-01", "v": 5}, {"c": None, "m": "2026-03", "v": 1},
    ]
    headers, out, bold = pivot_rows(rows, ["c"], "m", "v", True)
    assert headers == ["c", "2026-01", "2026-02", "2026-03", "Total"]
    assert out[0] == ["C1", 3.0, 2.0, None, 5.0] and out[2] == ["", None, None, 1.0, 1.0]
    assert out[-1] == ["Total", 8.0, 2.0, 1.0, 11.0] and bold == [False, False, False, True]
    headers, out, bold = pivot_rows(rows, ["c", "m"], None, "v", False)
    assert headers == ["c", "m", "v"] and out[0] == ["C1", "", None] and bold[0] is True and out[1] == ["", "2026-01", 3.0]
    assert not any(r[0] == "Total" for r in out)     # non-additive: no grand total


@pytest.mark.asyncio
async def test_sections_xlsx_and_markdown(env) -> None:  # noqa: F811
    registry, screens, view = env
    base = {"entity": "sales", "rows": [{"field": "company.NAME"}, {"field": "sales.SOLD_ON", "bucket": "month"}], "measures": [{"field": "sales.REVENUE", "agg": "sum"}], "filterRows": [], "sort": [], "limit": 100, "calculated": []}
    items = [
        ("Revenue table", {**base, "view": "table", "period": {"field": "sales.SOLD_ON", "preset": "custom", "range": {"dateFrom": "2026-01-01", "dateTo": "2026-12-31"}}}),
        ("Revenue pivot", {**base, "view": "pivot", "pivot": {"measure": "sum(sales.REVENUE)"}}),
        ("Revenue chart", {**base, "view": "bar"}),
        ("Revenue KPI", {**base, "view": "kpi"}),
        ("Sales detail", {**base, "view": "detail", "rows": [{"field": "sales.ID"}, {"field": "sales.REVENUE"}]}),
        ("Broken", {**base, "view": "table", "rows": [{"field": "sales.NOPE"}]}),
    ]
    sections = await build_sections(registry, screens, view, items)
    by = {s.title: s for s in sections}
    assert by["Revenue table"].headers == ["NAME", "SOLD_ON (month)", "sum REVENUE"]
    assert by["Revenue table"].rows[0][:2] == ["Corp one", "2026-01"] and "period 2026-01-01 → 2026-12-31" in (by["Revenue table"].note or "")
    assert by["Revenue pivot"].headers[-1] == "Total" and by["Revenue pivot"].rows[-1][0] == "Total" and by["Revenue pivot"].bold[-1]
    assert by["Revenue chart"].headers[0] == "NAME" and len(by["Revenue chart"].headers) == 5     # x + 4 months
    assert by["Revenue KPI"].rows == [[259.0]]
    assert by["Sales detail"].headers == ["ID", "REVENUE"] and len(by["Sales detail"].rows) == 6
    assert by["Broken"].kind == "error" and "unknown field" in by["Broken"].rows[0][0]
    # xlsx: one sheet per section, the headers bold, the note italic
    wb = load_workbook(io.BytesIO(sections_to_xlsx(sections)))
    assert wb.sheetnames[:2] == ["Revenue table", "Revenue pivot"]
    ws = wb["Revenue table"]
    assert ws.cell(row=1, column=1).font.italic and ws.cell(row=2, column=1).font.bold and ws.cell(row=3, column=1).value == "Corp one"
    md = sections_to_markdown(sections, "Sales report")
    assert md.startswith("# Sales report") and "## Revenue pivot" in md and "| **Total** |" in md and "| --- |" in md
