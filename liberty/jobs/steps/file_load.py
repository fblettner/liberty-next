"""``file_load`` executor — a spreadsheet / CSV file into a table, replayable by partition.

What the grid's Excel import is not: a load that runs without a person at the keyboard (a
schedule, a watched folder, a chain of steps) on volumes a facts table takes, that can be
replayed for one period without touching the others, and that leaves a trace.

The step (``jobs.toml``)::

    [[jobs.steps]]
    type = "file_load"
    name = "load sales month"
    path = "/data/in/sales-${params.period}.xlsx"     # ${params.x} from the run parameters, ${run.id}
    sheet = "Sales"                                     # name or 1-based index; default: the first
    header_row = 1
    target = { connector = "sales", table = "fact_sales" }
    columns = { article = "Item", qty = "Quantity", revenue = "#5" }   # target ← header, or #position
    constants = { period = "${params.period}", company = "${params.company}", batch_id = "${run.id}" }
    load_mode = "replace_partition"                     # or "append"
    partition_by = ["period", "company"]
    on_error = "fail"                                   # or "skip" (the row is counted as rejected)

* The file is streamed (openpyxl read-only for ``.xlsx`` / ``.xlsm``, the csv module for
  ``.csv`` / ``.txt``), never loaded whole.
* Target columns not in ``columns`` match a header of the same name (case, spaces and
  underscores ignored); ``constants`` fill the rest; whatever is left stays NULL.
* Values are coerced to the target column's type (introspected): integers, decimals, dates
  (ISO, ``dd/mm/yyyy``, Excel dates), booleans, text. A bad cell fails the step, or skips the
  row with ``on_error = "skip"``; the first ten rejects are quoted in the run log.
* ``replace_partition``: the partitions present in the load (from the constants, else the
  distinct values found in the file) are deleted and the rows inserted **in one transaction**,
  so a failed load leaves the previous data in place and a replay of a month replaces that
  month only. ``append`` just inserts.
"""

from __future__ import annotations

import csv
import logging
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import Boolean, Date, DateTime, Integer, Numeric, String, Text, text
from sqlalchemy.ext.asyncio import AsyncEngine
from sqlalchemy.types import TypeEngine

from liberty.connectors import ConnectorRegistry, UnknownConnectorError
from liberty.connectors.sql import SQLConnector
from liberty.jobs.coercion import ColumnInfo
from liberty.jobs.schema import LoadMode, Step, StepType
from liberty.jobs.steps.base import RunContext, StepFailed, StepResult

_log = logging.getLogger(__name__)

# Every placeholder carries a dot on purpose: the config loader replaces any ``${NAME}`` with an
# environment variable (empty when unset) before the job is parsed, and ``\w+`` stops at a dot.
_TEMPLATE = re.compile(r"\$\{(params\.[A-Za-z_][A-Za-z0-9_]*|run\.id|job\.id)\}")
_POSITION = re.compile(r"^#(\d+)$")
_SPREADSHEET = {".xlsx", ".xlsm"}
_DELIMITED = {".csv", ".txt", ".tsv"}
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d.%m.%Y", "%Y%m%d")
_DATETIME_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M")
_TRUE = {"1", "true", "yes", "y", "t", "oui", "o"}
_FALSE = {"0", "false", "no", "n", "f", "non", ""}
_MAX_QUOTED_REJECTS = 10


class CellError(ValueError):
    """A value the target column cannot take."""


def render_template(value: str, *, params: dict[str, Any], run_id: str, job_id: str) -> str:
    """``${params.x}`` / ``${run.id}`` / ``${job.id}`` substitution; an unknown parameter fails
    the step rather than silently loading into a wrong partition."""
    def sub(m: re.Match[str]) -> str:
        key = m.group(1)
        if key == "run.id":
            return run_id
        if key == "job.id":
            return job_id
        name = key[len("params."):]
        if name not in params or params[name] is None:
            raise StepFailed(f"file_load: parameter {name!r} is not set (needed by {value!r})")
        return str(params[name])
    return _TEMPLATE.sub(sub, value)


def normalize_header(h: Any) -> str:
    """Header / column names compare without case, spaces, underscores or dashes."""
    return re.sub(r"[\s_\-]+", "", str(h or "")).lower()


# --------------------------------------------------------------------------- #
# readers — yield (row_number, [cells…]) lazily, header first
# --------------------------------------------------------------------------- #


def _read_spreadsheet(path: Path, *, sheet: str | int | None) -> Iterator[tuple[int, list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover — openpyxl is a runtime dependency
        raise StepFailed("file_load: openpyxl is not installed") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None or sheet == "":
            ws = wb.worksheets[0]
        elif isinstance(sheet, int) or (isinstance(sheet, str) and sheet.isdigit()):
            idx = int(sheet) - 1
            if idx < 0 or idx >= len(wb.worksheets):
                raise StepFailed(f"file_load: {path.name} has no sheet #{sheet} ({len(wb.worksheets)} sheet(s))")
            ws = wb.worksheets[idx]
        else:
            if sheet not in wb.sheetnames:
                raise StepFailed(f"file_load: {path.name} has no sheet {sheet!r} (sheets: {', '.join(wb.sheetnames)})")
            ws = wb[sheet]
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            yield i, list(row)
    finally:
        wb.close()


def _read_delimited(path: Path, *, delimiter: str | None, encoding: str) -> Iterator[tuple[int, list[Any]]]:
    with path.open("r", encoding=encoding, newline="") as fh:
        if delimiter:
            reader = csv.reader(fh, delimiter=delimiter)
        else:
            sample = fh.read(8192)
            fh.seek(0)
            try:
                sniffed = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                sniffed = csv.excel  # type: ignore[assignment]
            reader = csv.reader(fh, dialect=sniffed)
        for i, row in enumerate(reader, start=1):
            yield i, list(row)


def open_reader(path: Path, *, sheet: str | int | None, delimiter: str | None, encoding: str) -> Iterator[tuple[int, list[Any]]]:
    ext = path.suffix.lower()
    if ext in _SPREADSHEET:
        return _read_spreadsheet(path, sheet=sheet)
    if ext in _DELIMITED:
        return _read_delimited(path, delimiter=delimiter, encoding=encoding)
    raise StepFailed(f"file_load: {path.name}: unsupported file type {ext or '(none)'} (xlsx, xlsm, csv, txt, tsv)")


# --------------------------------------------------------------------------- #
# coercion to the target column's type
# --------------------------------------------------------------------------- #


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def coerce_cell(value: Any, col_type: TypeEngine) -> Any:
    """A file cell as the target column takes it. Spreadsheets hand typed cells (numbers, dates),
    CSV hands text — both end up the same. Empty → NULL."""
    if _blank(value):
        return None
    if isinstance(col_type, Boolean):
        if isinstance(value, bool):
            return value
        s = str(value).strip().lower()
        if s in _TRUE:
            return True
        if s in _FALSE:
            return False
        raise CellError(f"not a boolean: {value!r}")
    if isinstance(col_type, Integer):
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            if value != int(value):
                raise CellError(f"not a whole number: {value!r}")
            return int(value)
        s = str(value).strip().replace(" ", "").replace(" ", "")
        try:
            d = Decimal(s.replace(",", "."))
        except InvalidOperation as exc:
            raise CellError(f"not a whole number: {value!r}") from exc
        if d != d.to_integral_value():
            raise CellError(f"not a whole number: {value!r}")
        return int(d)
    if isinstance(col_type, Numeric):  # Numeric covers Float / DECIMAL
        if isinstance(value, bool):
            raise CellError(f"not a number: {value!r}")
        if isinstance(value, (int, float, Decimal)):
            return value
        s = str(value).strip().replace(" ", "").replace(" ", "")
        if s.count(",") == 1 and "." not in s:
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
        try:
            return Decimal(s)
        except InvalidOperation as exc:
            raise CellError(f"not a number: {value!r}") from exc
    if isinstance(col_type, DateTime):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time())
        s = str(value).strip()
        for fmt in _DATETIME_FORMATS + _DATE_FORMATS:
            try:
                return datetime.strptime(s, fmt)
            except ValueError:
                continue
        raise CellError(f"not a date/time: {value!r}")
    if isinstance(col_type, Date):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        s = str(value).strip()
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        for fmt in _DATETIME_FORMATS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        raise CellError(f"not a date: {value!r}")
    if isinstance(col_type, (String, Text)):
        if isinstance(value, float) and value == int(value):
            return str(int(value))     # a code typed as a number by the spreadsheet ("00050" → 50.0)
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return str(value).strip()
    return value


# --------------------------------------------------------------------------- #
# the plan: which target column takes which file column / constant
# --------------------------------------------------------------------------- #


@dataclass(slots=True)
class LoadPlan:
    columns: list[ColumnInfo]
    """target columns in the table's order"""
    from_file: dict[str, int]
    """target column → index in the file row"""
    constants: dict[str, Any]
    """target column → coerced constant"""

    @property
    def loaded(self) -> list[str]:
        return [c.name for c in self.columns if c.name in self.from_file or c.name in self.constants]


def build_plan(
    header: list[Any], columns: list[ColumnInfo], *, mapping: dict[str, str], constants: dict[str, Any], path_name: str,
) -> LoadPlan:
    by_target = {c.name.lower(): c for c in columns}
    by_header: dict[str, int] = {}
    for i, h in enumerate(header):
        key = normalize_header(h)
        if key and key not in by_header:
            by_header[key] = i
    from_file: dict[str, int] = {}
    for target, source in mapping.items():
        col = by_target.get(target.lower())
        if col is None:
            raise StepFailed(f"file_load: column {target!r} is not in the target table")
        m = _POSITION.match(source.strip())
        if m:
            pos = int(m.group(1)) - 1
            if pos < 0 or pos >= len(header):
                raise StepFailed(f"file_load: {path_name}: column position {source} is outside the header ({len(header)} column(s))")
            from_file[col.name] = pos
        else:
            idx = by_header.get(normalize_header(source))
            if idx is None:
                raise StepFailed(f"file_load: {path_name}: header {source!r} not found (headers: {', '.join(str(h) for h in header if h is not None)})")
            from_file[col.name] = idx
    consts: dict[str, Any] = {}
    for target, value in constants.items():
        col = by_target.get(target.lower())
        if col is None:
            raise StepFailed(f"file_load: constant {target!r} is not a column of the target table")
        try:
            consts[col.name] = coerce_cell(value, col.type)
        except CellError as exc:
            raise StepFailed(f"file_load: constant {target!r}: {exc}") from exc
    for c in columns:   # same-name match for the rest
        if c.name in from_file or c.name in consts:
            continue
        idx = by_header.get(normalize_header(c.name))
        if idx is not None:
            from_file[c.name] = idx
    if not from_file:
        raise StepFailed(f"file_load: {path_name}: no file column matches the target table (map them with `columns`)")
    return LoadPlan(columns=columns, from_file=from_file, constants=consts)


def _row_values(plan: LoadPlan, cells: list[Any]) -> dict[str, Any]:
    out: dict[str, Any] = dict(plan.constants)
    types = {c.name: c.type for c in plan.columns}
    for name, idx in plan.from_file.items():
        v = cells[idx] if idx < len(cells) else None
        try:
            out[name] = coerce_cell(v, types[name])
        except CellError as exc:
            raise CellError(f"{name}: {exc}") from exc
    return out


def _adapt(row: dict[str, Any], dialect: str) -> dict[str, Any]:
    """Driver quirks on the bind side: sqlite3 takes neither Decimal nor date objects through
    a textual statement — numbers go as float, dates as ISO text (what the column holds anyway)."""
    if dialect != "sqlite":
        return row
    out = dict(row)
    for k, v in out.items():
        if isinstance(v, Decimal):
            out[k] = float(v)
        elif isinstance(v, (date, datetime)):
            out[k] = v.isoformat()
    return out


# --------------------------------------------------------------------------- #
# executor
# --------------------------------------------------------------------------- #


class FileLoadExecutor:
    """Executes one ``file_load`` step. Stateless."""

    def __init__(self, connectors: ConnectorRegistry) -> None:
        self._connectors = connectors

    async def execute(self, step: Step, ctx: RunContext) -> StepResult:
        if step.type is not StepType.FILE_LOAD:
            raise StepFailed(f"FileLoadExecutor received a step of type {step.type.value!r}")
        assert step.target is not None and step.path is not None
        params = dict(step.op_kwargs)   # the runner merged job.params + the per-fire override here
        render = lambda s: render_template(s, params=params, run_id=ctx.run_id, job_id=ctx.job_id)  # noqa: E731
        path = Path(render(step.path)).expanduser()
        if not path.is_file():
            raise StepFailed(f"file_load: file not found: {path}")
        constants = {k: (render(v) if isinstance(v, str) else v) for k, v in step.constants.items()}
        mode = step.load_mode or LoadMode.APPEND
        partition_by = [p for p in step.partition_by]

        try:
            tgt = self._connectors.get(step.target.connector)
        except UnknownConnectorError as exc:
            raise StepFailed(f"file_load: unknown target connector {step.target.connector!r}") from exc
        if not isinstance(tgt, SQLConnector):
            raise StepFailed(f"file_load: target {step.target.connector!r} is not a SQL connector")
        engine: AsyncEngine = self._connectors.pools.engine(tgt.pool_name)
        from liberty.jobs.steps.sql_copy import _introspect_columns, _qualified
        columns = await _introspect_columns(engine, schema=step.target.schema_, table=step.target.table)
        if not columns:
            raise StepFailed(f"file_load: target table {step.target.schema_ or ''}.{step.target.table} has no columns (missing?)")
        table_sql = _qualified(step.target.schema_, step.target.table)
        by_name = {c.name.lower(): c.name for c in columns}
        for p in partition_by:
            if p.lower() not in by_name:
                raise StepFailed(f"file_load: partition column {p!r} is not in the target table")
        partition_by = [by_name[p.lower()] for p in partition_by]

        _log.info(
            "nomaflow.file_load start run=%s step=%r file=%s → %s.%s mode=%s partition_by=%s",
            ctx.run_id, step.name, path, tgt.name, step.target.table, mode.value, partition_by or None,
        )

        # pass 1 — header + plan (+ the partitions present, when a key isn't a constant)
        reader = open_reader(path, sheet=step.sheet, delimiter=step.delimiter, encoding=step.encoding)
        header = self._header(reader, step.header_row, path.name)
        plan = build_plan(header, columns, mapping=step.columns, constants=constants, path_name=path.name)
        missing_keys = [p for p in partition_by if p not in plan.loaded]
        if missing_keys:
            raise StepFailed(f"file_load: partition column(s) {', '.join(missing_keys)} come neither from the file nor from `constants`")
        partitions: set[tuple[Any, ...]] = set()
        if mode is LoadMode.REPLACE_PARTITION:
            if all(p in plan.constants for p in partition_by):
                partitions.add(tuple(plan.constants[p] for p in partition_by))
            else:
                for _n, cells in reader:
                    if self._empty(cells):
                        continue
                    try:
                        vals = _row_values(plan, cells)
                    except CellError:
                        continue   # pass 2 reports it
                    partitions.add(tuple(vals[p] for p in partition_by))

        # pass 2 — one transaction: delete the partitions, insert every batch
        names = plan.loaded
        insert_sql = text(f"INSERT INTO {table_sql} ({', '.join(_q(n) for n in names)}) VALUES ({', '.join(':' + n for n in names)})")
        read = inserted = rejected = 0
        rejects: list[str] = []
        reader = open_reader(path, sheet=step.sheet, delimiter=step.delimiter, encoding=step.encoding)
        self._header(reader, step.header_row, path.name)
        async with engine.connect() as conn:
            async with conn.begin():
                deleted = 0
                for key in sorted(partitions, key=lambda k: tuple(str(x) for x in k)):
                    where = " AND ".join(f"{_q(p)} = :p{i}" for i, p in enumerate(partition_by))
                    res = await conn.execute(text(f"DELETE FROM {table_sql} WHERE {where}"), _adapt({f"p{i}": v for i, v in enumerate(key)}, engine.dialect.name))
                    deleted += res.rowcount or 0
                batch: list[dict[str, Any]] = []
                for n, cells in reader:
                    if self._empty(cells):
                        continue
                    read += 1
                    try:
                        batch.append(_adapt(_row_values(plan, cells), engine.dialect.name))
                    except CellError as exc:
                        rejected += 1
                        if len(rejects) < _MAX_QUOTED_REJECTS:
                            rejects.append(f"row {n}: {exc}")
                        if step.on_error == "fail":
                            raise StepFailed(f"file_load: {path.name} row {n}: {exc}") from exc
                        continue
                    if len(batch) >= step.batch_size:
                        await conn.execute(insert_sql, batch)
                        inserted += len(batch)
                        batch = []
                if batch:
                    await conn.execute(insert_sql, batch)
                    inserted += len(batch)
        for line in rejects:
            _log.warning("nomaflow.file_load run=%s step=%r rejected %s", ctx.run_id, step.name, line)
        _log.info(
            "nomaflow.file_load done run=%s step=%r read=%d inserted=%d rejected=%d partitions=%d deleted=%d",
            ctx.run_id, step.name, read, inserted, rejected, len(partitions), deleted,
        )
        return StepResult(
            rows_affected=inserted,
            log_excerpt=f"{path.name}: {read} row(s) read, {inserted} inserted, {rejected} rejected"
            + (f", {len(partitions)} partition(s) replaced ({deleted} row(s) deleted)" if mode is LoadMode.REPLACE_PARTITION else ""),
            extras={
                "file": str(path), "read": read, "inserted": inserted, "rejected": rejected, "deleted": deleted,
                "partitions": [list(k) for k in sorted(partitions, key=lambda k: tuple(str(x) for x in k))],
                "rejects": rejects,
            },
        )

    @staticmethod
    def _header(reader: Iterator[tuple[int, list[Any]]], header_row: int, name: str) -> list[Any]:
        for n, cells in reader:
            if n == header_row:
                if not any(not _blank(c) for c in cells):
                    raise StepFailed(f"file_load: {name}: header row {header_row} is empty")
                return cells
        raise StepFailed(f"file_load: {name}: no header row {header_row} (file too short)")

    @staticmethod
    def _empty(cells: Iterable[Any]) -> bool:
        return all(_blank(c) for c in cells)


def _q(name: str) -> str:
    return f'"{name}"'
